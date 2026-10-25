"""Tes logika app.py tanpa memuat YOLO: predict_raw dan get_names diganti stub.

Jalankan: pytest -q
"""
import io
import sys
from pathlib import Path

import cv2
import numpy as np
import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import app  # noqa: E402

NAMES = {0: "apple_fresh", 1: "apple_rotten", 2: "banana_fresh", 3: "banana_rotten"}


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(app, "DB_PATH", str(tmp_path / "t.db"))
    monkeypatch.setattr(app, "RECENT_DIR", tmp_path / "recent")
    monkeypatch.setattr(app, "REVIEW_DIR", tmp_path / "review")
    monkeypatch.setattr(app, "get_names", lambda: NAMES)
    app.init_db()
    return tmp_path


def png(w=64, h=48):
    ok, buf = cv2.imencode(".png", np.zeros((h, w, 3), np.uint8))
    return buf.tobytes()


def stub(monkeypatch, dets):
    def fake(img, conf, iou, class_ids, annotate=True):
        keep = [d for d in dets if d[1] >= conf and (class_ids is None or d[0] in class_ids)]
        return keep, (img if annotate else None)
    monkeypatch.setattr(app, "predict_raw", fake)


def detect(name="a.png", data=None, conf=0.5, classes=None):
    return app.run_detection(data or png(), name, "upload", conf, 0.4, classes)


def test_fresh_gives_price(env, monkeypatch):
    stub(monkeypatch, [(0, 0.9, [1, 2, 30, 40])])
    r = detect()
    assert r["status"] == "fresh" and r["total_price"] == 8500
    assert r["annotated_rgb"].shape == (48, 64, 3)


def test_rotten_blocks_price(env, monkeypatch):
    stub(monkeypatch, [(0, 0.9, [1, 2, 3, 4]), (3, 0.8, [5, 6, 7, 8])])
    r = detect()
    assert r["status"] == "rotten" and r["total_price"] == 0


def test_no_object_and_conf_filter(env, monkeypatch):
    stub(monkeypatch, [(0, 0.4, [1, 2, 3, 4])])
    assert detect(conf=0.5)["status"] == "none"


def test_class_filter_exact_and_by_fruit(env, monkeypatch):
    stub(monkeypatch, [(0, 0.9, [1, 2, 3, 4]), (1, 0.9, [5, 6, 7, 8]), (2, 0.9, [9, 10, 11, 12])])
    assert [o["fruit"] for o in detect(classes=["banana"])["objects"]] == ["banana"]
    assert [o["label"] for o in detect(classes=["apple_rotten"])["objects"]] == ["apple_rotten"]
    with pytest.raises(app.AppError):
        detect(classes=["durian"])
    with pytest.raises(app.AppError):
        detect(classes=123)


def test_invalid_empty_and_oversize_file(env):
    with pytest.raises(app.AppError):
        detect(data=b"bukan gambar")
    with pytest.raises(app.AppError):
        app.run_detection(b"", "x.png", "upload", 0.5, 0.4)
    with pytest.raises(app.AppError, match="10 MB"):
        app.run_detection(b"x" * (app.MAX_IMAGE_BYTES + 1), "x.png", "upload", 0.5, 0.4)


def test_history_sort_whitelist(env, monkeypatch):
    stub(monkeypatch, [(0, 0.9, [1, 2, 3, 4])])
    detect()
    assert app.history("n_objects", "asc")["total"] == 1
    with pytest.raises(app.AppError):
        app.history("id;DROP TABLE scans")
    with pytest.raises(app.AppError):
        app.history(order="sideways")


def test_feedback_once_per_scan(env, monkeypatch):
    stub(monkeypatch, [(1, 0.77, [1, 2, 3, 4])])
    sid = detect()["scan_id"]
    items = [{"fruit": "apple", "confidence": 0.77}]
    assert app.save_feedback(sid, "Rotten (Afkir)", items) == "Dibuang / Retur Supplier"
    with pytest.raises(app.AppError, match="sudah punya keputusan"):
        app.save_feedback(sid, "False Positive (Koreksi Staf)", items)
    assert "77.00%" in app.waste_log_csv().decode() and "Dibuang / Retur Supplier" in app.waste_log_csv().decode()
    with pytest.raises(app.AppError, match="tidak ditemukan"):
        app.save_feedback(999, "Rotten (Afkir)", items)
    with pytest.raises(app.AppError):
        app.save_feedback(sid, "Status ngawur", items)


def test_low_conf_goes_to_review_queue(env, monkeypatch):
    stub(monkeypatch, [(0, 0.55, [1, 2, 3, 4])])
    sid = detect(conf=0.25)["scan_id"]
    assert (app.REVIEW_DIR / f"{sid}_lowconf.jpg").exists()


def test_stats_and_exports(env, monkeypatch):
    stub(monkeypatch, [(0, 0.95, [1, 2, 3, 4]), (3, 0.5, [5, 6, 7, 8])])
    detect()
    s = app.stats()
    assert s["total_objects"] == 2 and s["conditions"] == {"fresh": 1, "rotten": 1}
    assert sum(b["count"] for b in s["confidence_hist"]) == 2
    rows = app.det_query()
    assert app.export_bytes("csv", rows)[0].decode().count("\n") == 3
    assert len(app.det_query(label="banana")) == 1
    assert app.export_bytes("xlsx", rows)[0][:2] == b"PK"


def test_feedback_kept_when_review_copy_fails(env, monkeypatch):
    stub(monkeypatch, [(0, 0.9, [1, 2, 3, 4])])
    sid = detect()["scan_id"]
    monkeypatch.setattr(app.shutil, "copy", lambda *a, **k: (_ for _ in ()).throw(OSError("disk penuh")))
    items = [{"fruit": "apple", "confidence": 0.9}]
    app.save_feedback(sid, "False Negative (Koreksi Staf)", items)
    assert len(app.waste_log()) == 1


def test_exports_neutralize_spreadsheet_formulas(env, monkeypatch):
    stub(monkeypatch, [(0, 0.9, [1, 2, 3, 4])])
    r = detect(name='=HYPERLINK("http://x","klik").png')
    rows = app.det_query()
    assert "'=HYPERLINK" in app.export_bytes("csv", rows)[0].decode()
    ws = load_workbook(io.BytesIO(app.export_bytes("xlsx", rows)[0])).active
    cell = ws.cell(row=2, column=3)
    assert cell.value.startswith("=HYPERLINK") and cell.data_type == "s"
    app.save_feedback(r["scan_id"], "Rotten (Afkir)", [{"fruit": "=cmd|' /C calc'!A0", "confidence": 0.9}])
    assert "'=cmd" in app.waste_log_csv().decode()


def test_sales_numbering_and_today(env):
    assert app.next_receipt_no() == 1
    app.save_sale(17000, 2, 0, 20000, [{"fruit": "apple", "price": 8500, "qty": 2}])
    assert app.next_receipt_no() == 2
    assert app.today_sales() == {"n": 1, "sum": 17000}


def test_receipt_text_has_total(env):
    t = {"sub": 20000, "d": 10, "disc": 2000, "total": 18000, "cash": 20000, "change": 2000, "short": False}
    txt = app.receipt_text(1, "2026-10-25 10:00:00", [{"fruit": "apple", "price": 10000, "qty": 2}], {"banana": 1}, t)
    assert "Rp 18.000" in txt and "Diskon 10%" in txt and "Pisang x1" in txt
